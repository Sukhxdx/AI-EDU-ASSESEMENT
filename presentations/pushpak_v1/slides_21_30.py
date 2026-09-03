"""Slides 21–30 — sensors, modes, safety, mass, BOM, analysis plans."""

from pathlib import Path

from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

import kit as k

A = Path(__file__).parent / "assets"


def _base(prs, kicker, title, num, notes, subtitle=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    k.set_bg(s)
    k.header(s, kicker, title, subtitle)
    k.footer(s, num)
    k.notes(s, notes)
    return s


def s21_sensors(prs):
    s = _base(
        prs, "20  /  SENSORS", "Sensor Suite", 21,
        "Every sensor exists because a later plot needs it. Pitot for tilt schedule and CL. Mag angle for the two shafts. V/I for the power experiment.\n"
        "Compass + GPS are navigation and RTL — not the science.\n"
        "Absolute magnetic angle is preferred over servo-potentiometer because the servo horn is not the shaft.",
    )
    sensors = [
        ("PITOT", "Calibrated airspeed", "Transition gate  ·  CL reconstruction"),
        ("GPS", "Position / groundspeed", "Geofence  ·  RTL  ·  track"),
        ("COMPASS", "Heading", "Yaw reference in hover & cruise"),
        ("MAG ANGLE ×2", "Shaft tilt", "Closed-loop tilt  ·  jam detect"),
        ("IMU (FC)", "Rates + accel", "Att. loop  ·  n_z during transition"),
        ("BARO (FC)", "Altitude / Vz", "Climb hold in QHOVER / QLOITER"),
        ("V / I BUS", "Electrical power", "Primary experiment observable"),
        ("I_ACTUATOR", "Servo / ESC current", "Jam, stall, ESC health"),
    ]
    for i, (t, a, b) in enumerate(sensors):
        x = 0.42 + (i % 4) * 3.20
        y = 1.22 + (i // 4) * 2.90
        k.card(s, x, y, 3.05, 2.70, k.CYAN)
        k.textbox(s, x + 0.16, y + 0.18, 2.75, 0.32, t, 13, k.CYAN, True, k.MONO)
        k.textbox(s, x + 0.16, y + 0.58, 2.75, 0.70, a, 16, k.WHITE, True)
        k.textbox(s, x + 0.16, y + 1.40, 2.75, 1.05, b, 13, k.OFF)


def s22_modes(prs):
    s = _base(
        prs, "21  /  MODES", "Flight Modes", 22,
        "ArduPilot VTOL vocabulary. First flights stay in QSTABILIZE / QHOVER. Transition is a gated climb into a forward mode, not a leap into AUTO.\n"
        "FBWA or CRUISE is the wing-borne science mode. AUTO is last.\n"
        "A mode that cannot revert to hover in one switch is not a student mode.",
    )
    modes = [
        ("QSTABILIZE", "01", "Manual throttle, attitude stabilize. Tilt locked at 0°. First tether hover."),
        ("QHOVER", "02", "Alt-hold hover. Used for CG check, tilt-servo health, and pad ops."),
        ("QLOITER", "03", "Position hold. Precision take-off and landing once GPS is trusted."),
        ("TRANSITION", "04", "Airspeed-scheduled tilt 0→75°. Collective reduced as wing unloads rotors."),
        ("FBWA / CRUISE", "05", "Wing-borne science window. Power logging at 20 m/s. Tilt held at 75°."),
        ("RTL / QRTL", "06", "Failsafe home. Tilt back toward hover, then vertical descent on pad."),
    ]
    for i, (t, n, b) in enumerate(modes):
        x = 0.42 + (i % 3) * 4.16
        y = 1.22 + (i // 3) * 2.90
        k.card(s, x, y, 3.98, 2.70, k.CYAN)
        k.textbox(s, x + 0.18, y + 0.16, 3.6, 0.28, n, 12, k.CYAN, True, k.MONO)
        k.textbox(s, x + 0.18, y + 0.48, 3.6, 0.50, t, 18, k.WHITE, True)
        k.textbox(s, x + 0.18, y + 1.15, 3.6, 1.30, b, 14, k.OFF)


def s23_transition(prs):
    s = _base(
        prs, "22  /  TRANSITION", "Transition Strategy", 23,
        "Transition is a schedule, not a stunt. Gates are airspeed, climb rate, and tilt-rate limit.\n"
        "Front/rear split is a pitch trim tool — if the nose drops, lag the front shaft.\n"
        "Abort is always available: freeze tilt, add collective, settle to QHOVER.",
    )
    k.process_bar(
        s, 0.42, 1.22, 12.48, 0.50,
        ["Pad hover", "Nose into wind", "Build V", "Tilt ramp", "Wing hold", "Cruise log"],
        active=5,
    )
    steps = [
        ("T+0", "QHOVER 1.5–2 m AGL. Confirm mag-angle 0°, bus healthy, pitot alive."),
        ("T+1", "Command a small positive pitch / forward cyclic. Groundspeed rising, Vz ≈ 0."),
        ("T+2", "At V1 (to be set in taxi tests, expected 8–12 m/s) start tilt ramp, limited °/s."),
        ("T+3", "As CL builds, lower collective so that altitude does not balloon."),
        ("T+4", "Hold 75° once V ≈ 20 m/s and n_z is quiet. Open the power-log window."),
        ("T+5", "Abort line: airspeed decay, tilt jam, or Vz > band → reverse tilt, QHOVER."),
    ]
    for i, (t, b) in enumerate(steps):
        y = 1.90 + i * 0.80
        k.card(s, 0.42, y, 12.48, 0.72, k.CYAN if i < 5 else k.AMBER)
        k.textbox(s, 0.58, y + 0.16, 1.3, 0.42, t, 16, k.CYAN if i < 5 else k.AMBER, True, k.MONO)
        k.textbox(s, 2.05, y + 0.16, 10.5, 0.44, b, 14, k.OFF)


def s24_safety(prs):
    s = _base(
        prs, "23  /  SAFETY", "Safety Features", 24,
        "Mechanical stops are not optional. A servo that can drive through 75° will put a blade into the wing.\n"
        "Electrical: current limits, low-voltage RTL, lost-link RTL.\n"
        "Operational: tether for first hovers, geofence, two-person crew (pilot + caller).",
    )
    data = [
        ["Layer", "Feature", "Intent"],
        ["Mechanical", "Hard stops at 0° and 75°", "Blade / wing collision impossible via software fault"],
        ["Mechanical", "Shaft bearings carry rotor loads", "Servo is a commander, not a structure"],
        ["Sense", "Dual mag-angle disagreement", "Inhibit further tilt; hold last safe angle"],
        ["Sense", "Pitot fail detect", "Inhibit transition; remain in VTOL modes"],
        ["Electrical", "Bus undervoltage", "QRTL / land"],
        ["Electrical", "Actuator overcurrent", "Freeze tilt; alert GCS"],
        ["FC", "Lost GCS / geofence", "QRTL"],
        ["FC", "Companion watchdog", "ESP32 death does not kill motors"],
        ["Ops", "Tethered first hover", "Contain first-flight energy"],
        ["Ops", "Prop clearance 50–75 mm", "Flex and downwash margin"],
    ]
    k.add_table(s, data, 0.42, 1.18, 12.48, 5.95, 12)


def s25_mass(prs):
    s = _base(
        prs, "24  /  MASS", "Mass Budget", 25,
        "Read the bar chart first: battery and payload dominate. Motors are heavy because hover margin is non-negotiable.\n"
        "Empty target 3.5 kg leaves 1.0 kg payload inside a 4.5 kg design point. Estimated roll-up today is 4.4–4.9 kg — we are on the edge and we know it.\n"
        "If battery lands at 1.5 kg, something else must come out of structure or we fly a 4.9 kg aircraft and recompute hover.",
    )
    k.picture(s, A / "mass_budget.png", 0.30, 1.18, 7.3, 4.55)
    data = [
        ["Group", "Mass (kg)"],
        ["Payload", "1.00"],
        ["Motors", "0.80"],
        ["ESCs", "0.25"],
        ["Battery", "1.0–1.5"],
        ["Tilt system", "0.30"],
        ["Wing", "0.30"],
        ["Fuselage", "0.40"],
        ["Electronics", "0.20"],
        ["Landing gear", "0.15"],
        ["Estimated MTOW", "4.4–4.9"],
        ["Empty target", "≤ 3.50"],
        ["Design point", "4.50"],
    ]
    k.add_table(s, data, 7.70, 1.18, 5.20, 5.95, 11)


def s26_bom(prs):
    s = _base(
        prs, "25  /  BOM", "Component Selection Matrix", 26,
        "This is the buying list, not a catalogue dump. Six groups map to the mass budget.\n"
        "STS3215-C018 is a bus servo with stall torque suited to a 4.5 kg tilt shaft, not a hobby 9-gram.\n"
        "Pixhawk + ArduPilot is chosen for VTOL heritage, not because we cannot write a controller.",
        "Propulsion  ·  tilt  ·  flight control  ·  electronics  ·  battery  ·  airframe",
    )
    tables = [
        (0.42, 1.18, "PROPULSION", [
            ["Item", "Spec", "Qty"],
            ["Motor", "HGLRC MY4215 400KV", "4"],
            ["Propeller", "16-inch class / 406 mm", "4"],
            ["ESC", "6S, 0.25 kg set", "4"],
        ]),
        (4.55, 1.18, "TILT MECHANISM", [
            ["Item", "Spec", "Qty"],
            ["Servo", "STS3215-C018", "2"],
            ["Tilt shaft + mounts", "Front / rear axis", "2+4"],
            ["Bearings / stops", "4 per axis + hard stops", "8+"],
        ]),
        (8.68, 1.18, "FLIGHT CONTROL", [
            ["Item", "Spec", "Qty"],
            ["FC", "Pixhawk", "1"],
            ["Stack", "ArduPilot VTOL", "1"],
            ["Companion", "ESP32", "1"],
        ]),
        (0.42, 4.20, "ELECTRONICS", [
            ["Item", "Spec", "Qty"],
            ["Airspeed", "Pitot probe + sensor", "1"],
            ["Nav", "GPS + compass", "1+1"],
            ["Tilt sense", "Abs. magnetic angle", "2"],
        ]),
        (4.55, 4.20, "BATTERY SYSTEM", [
            ["Item", "Spec", "Qty"],
            ["Pack", "6S Li-ion ~20 Ah", "1"],
            ["Sense", "Bus V and I", "1"],
            ["Mass alloc.", "1.0–1.5 kg", "—"],
        ]),
        (8.68, 4.20, "AIRFRAME", [
            ["Item", "Spec", "Qty"],
            ["Wing", "1300 mm  NACA 4412", "1"],
            ["Winglets", "Tip treatments", "2"],
            ["Fuselage / gear", "1200–1300 mm / skids", "1"],
        ]),
    ]
    for x, y, title, data in tables:
        k.textbox(s, x, y, 3.9, 0.24, title, 10, k.CYAN, True, k.MONO)
        k.add_table(s, data, x, y + 0.26, 3.90, 2.45, 9)


def s27_cad(prs):
    s = _base(
        prs, "26  /  CAD", "CAD Development Plan", 27,
        "CAD is the configuration control system. Master assembly owns CG and clearance.\n"
        "Tilt axis is the first model to freeze — everything else hangs off those shafts.\n"
        "Weigh properties from CAD weekly against the mass budget; do not wait for the prototype.",
    )
    k.process_bar(s, 0.42, 1.22, 12.48, 0.48, ["Master", "Tilt axis", "Airframe", "Harness", "CG / MOI", "Drawings"])
    items = [
        ("WP-1  MASTER", "Top-down skeleton: wing, fuselage, two shafts, prop disks as clearance solids. Frozen interfaces."),
        ("WP-2  TILT", "Shaft, mounts, bearings, servo horn, linkage, stops, encoder target. Motion study 0–75°."),
        ("WP-3  AIRFRAME", "Wing skins/spars, winglets, fuselage bays for 6S pack, Pixhawk, payload 1 kg hard point."),
        ("WP-4  EQUIPMENT", "Motor/ESC/servo envelopes, pitot boom, GPS mast, antenna clear of props."),
        ("WP-5  MASS", "Assigned densities, live CG vs MAC and vs shaft plane. MOI for hover tuning."),
        ("WP-6  RELEASE", "Exploded views, 2D for 3D-print, laser, and composite templates. Configuration log."),
    ]
    for i, (t, b) in enumerate(items):
        x = 0.42 + (i % 3) * 4.16
        y = 1.90 + (i // 3) * 2.50
        k.card(s, x, y, 3.98, 2.32, k.CYAN)
        k.textbox(s, x + 0.16, y + 0.14, 3.66, 0.36, t, 13, k.CYAN, True, k.MONO)
        k.textbox(s, x + 0.16, y + 0.58, 3.66, 1.55, b, 13, k.OFF)


def s28_cfd(prs):
    s = _base(
        prs, "27  /  CFD", "CFD Validation Plan", 28,
        "CFD is not decoration. Three campaigns: isolated wing polar, hover downwash on the wing, and a transition snapshot.\n"
        "We will not claim full-aircraft accuracy in ground effect. We will claim trends and a CL check against 0.395.\n"
        "Power comparison remains a flight measurement; CFD sizes the wing and flags surprises.",
    )
    cases = [
        ("C1  WING POLAR", "Isolated wing + winglets. α sweep at Re ≈ 3.3×10⁵ (c=0.246 m, 20 m/s). Extract CL, CD, CM. Confirm CL=0.395 is a comfortable α."),
        ("C2  HOVER WASH", "Four 16-inch disks as actuator disks or MRF, wing at 0° tilt. Map download penalty and elevator blanking."),
        ("C3  CRUISE", "Full airframe, rotors at 75°, V=20 m/s. Residual rotor lift vs wing lift split. Nacelle / boom drag."),
        ("C4  TRANSITION", "One or two snapshots at 30–45° tilt. Identify pitch-up/pitch-down and a safe ramp rate."),
    ]
    for i, (t, b) in enumerate(cases):
        y = 1.22 + i * 1.18
        k.card(s, 0.42, y, 8.35, 1.08, k.CYAN)
        k.textbox(s, 0.60, y + 0.10, 8.0, 0.28, t, 13, k.CYAN, True, k.MONO)
        k.textbox(s, 0.60, y + 0.42, 8.0, 0.55, b, 13, k.OFF)
    k.card(s, 8.95, 1.22, 3.95, 5.80, k.AMBER)
    k.textbox(s, 9.12, 1.42, 3.6, 0.30, "SUCCESS METRICS", 12, k.AMBER, True, k.MONO)
    k.bullet_block(s, 9.05, 1.90, 3.7, 4.8, [
        "CL_req 0.395 inside linear range.",
        "Wing stall α > cruise α + 6°.",
        "Download increment quantified.",
        "No hidden pitch-up at 30–45°.",
        "Drag polar for power estimate.",
        "Mesh / residual report in annex.",
    ], 13)


def s29_structures(prs):
    s = _base(
        prs, "28  /  STRUCTURES", "Structural Analysis Plan", 29,
        "This is a 4.5 kg UAV, not a Part 23 aircraft. We still analyse the parts that kill people or the project: blades into wing, shaft yield, motor-mount fatigue, landing spike.\n"
        "Load cases are conservative factors on hover thrust and a 3-point landing.\n"
        "3D-printed mounts get extra scrutiny — they are the usual failure.",
    )
    data = [
        ["ID", "Load case", "Factor / condition", "Pass criterion"],
        ["LC-1", "Hover thrust on mounts", "2.25 kgf × 2.0 per motor", "No yield; bolt bearing OK"],
        ["LC-2", "Tilt-shaft torsion / bending", "Max motor torque + gyro", "Deflection does not eat clearance"],
        ["LC-3", "Wing +2.5 / −1.0 g", "Cruise 20 m/s, 70% lift", "Spar cap / foam shear OK"],
        ["LC-4", "Landing 3.5 g spike", "Skids, fuselage keel", "No FC tray crush"],
        ["LC-5", "Servo linkage stall", "STS3215 stall torque", "Horn / pin shear OK; stop takes load"],
        ["LC-6", "Prop burst / FOD", "Containment awareness", "Layout keeps blades off boom/wing"],
        ["LC-7", "Modal (shaft + wing)", "Hover RPM range", "No 1/rev coincidence"],
        ["LC-8", "CG / hard-point", "1 kg payload 10 g", "Insert pull-out margin"],
    ]
    k.add_table(s, data, 0.42, 1.18, 12.48, 5.95, 12)


def s30_mfg(prs):
    s = _base(
        prs, "29  /  MANUFACTURE", "Prototype Manufacturing Plan", 30,
        "Student-build rule: COTS propulsion, printed interfaces, simple wing. Do not open a composites factory for V1.\n"
        "Tilt shafts are the precision parts — buy tube, machine ends, or print jigs.\n"
        "First article is a fit-check airframe before we commit foam/composite skins.",
    )
    steps = [
        ("01", "PROPULSION", "Source MY4215, 16-in props, 6S ESCs. Thrust-stand characterisation before airframe lock."),
        ("02", "TILT AXIS", "Shaft, printed/CNC mounts, bearings, stops, STS3215, mag targets. Bench 0–75° cycle test."),
        ("03", "WING", "Hot-wire or CNC foam core, spars, NACA 4412 templates, bonded winglets. Mass ≤ 0.30 kg."),
        ("04", "FUSELAGE", "Pod for battery, Pixhawk, payload. Skids. Harness channels. Access panels."),
        ("05", "INTEGRATION", "Harness, pitot boom, GPS, encoder routing, CG weigh, control-surface or trim if any."),
        ("06", "QA", "Prop clearance 50–75 mm, tilt-stop inspection, insulation, photo log, mass report."),
    ]
    for i, (n, t, b) in enumerate(steps):
        x = 0.42 + (i % 3) * 4.16
        y = 1.22 + (i // 3) * 2.90
        k.card(s, x, y, 3.98, 2.70, k.CYAN)
        k.textbox(s, x + 0.16, y + 0.14, 3.66, 0.28, n, 12, k.CYAN, True, k.MONO)
        k.textbox(s, x + 0.16, y + 0.46, 3.66, 0.40, t, 16, k.WHITE, True)
        k.textbox(s, x + 0.16, y + 1.00, 3.66, 1.45, b, 13, k.OFF)
